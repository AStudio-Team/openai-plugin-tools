#!/usr/bin/env python3
"""Extract tables from a PDF and write them to an XLSX workbook."""

from __future__ import annotations

import argparse
import base64
import io
import json
import re
from pathlib import Path

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

MAX_PDF_BYTES = 20 * 1024 * 1024
MAX_PDF_PAGES = 200
MAX_TABLE_CELLS = 250_000
MAX_XLSX_BYTES = 10 * 1024 * 1024
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


class ConversionError(ValueError):
    """Raised when a PDF cannot be converted safely."""


def _sheet_name(page_number: int, table_number: int, used: set[str]) -> str:
    base = INVALID_SHEET_CHARS.sub("-", f"Page {page_number} Table {table_number}")[:31]
    candidate = base
    suffix = 2
    while candidate in used:
        tail = f" {suffix}"
        candidate = f"{base[:31-len(tail)]}{tail}"
        suffix += 1
    used.add(candidate)
    return candidate


def _safe_cell(value):
    """Keep untrusted PDF text from becoming an executable spreadsheet formula."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def convert_pdf_bytes(pdf_bytes: bytes) -> tuple[bytes, dict]:
    if not pdf_bytes.startswith(b"%PDF-"):
        raise ConversionError("Input is not a PDF file.")
    if len(pdf_bytes) > MAX_PDF_BYTES:
        raise ConversionError("PDF exceeds the 20 MB limit.")

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["PDF Table Extraction"])
    summary["A1"].font = Font(bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="322625")
    summary.append(["Page", "Table", "Rows", "Columns", "Worksheet"])
    for cell in summary[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="C0E3E5")

    used = {"Summary"}
    previews: list[dict] = []
    table_count = 0

    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            page_count = len(pdf.pages)
            if page_count > MAX_PDF_PAGES:
                raise ConversionError(f"PDF exceeds the {MAX_PDF_PAGES}-page limit.")
            total_cells = 0
            for page_number, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables() or []
                for table_number, table in enumerate(tables, start=1):
                    if not table or not any(any(cell not in (None, "") for cell in row) for row in table):
                        continue
                    table_count += 1
                    name = _sheet_name(page_number, table_number, used)
                    sheet = workbook.create_sheet(name)
                    width = max((len(row) for row in table), default=0)
                    total_cells += len(table) * width
                    if total_cells > MAX_TABLE_CELLS:
                        raise ConversionError(f"Extracted tables exceed the {MAX_TABLE_CELLS:,}-cell limit.")
                    for row in table:
                        normalized = [_safe_cell(value) for value in list(row)] + [None] * (width - len(row))
                        sheet.append(normalized)
                    if sheet.max_row:
                        for cell in sheet[1]:
                            cell.font = Font(bold=True)
                            cell.fill = PatternFill("solid", fgColor="C0E3E5")
                    summary.append([page_number, table_number, len(table), width, name])
                    previews.append({
                        "page": page_number,
                        "table": table_number,
                        "worksheet": name,
                        "rows": len(table),
                        "columns": width,
                        "preview": table[:5],
                    })
    except Exception as exc:
        raise ConversionError(f"Could not read PDF: {exc}") from exc

    if table_count == 0:
        raise ConversionError("No extractable tables were found in the PDF.")

    output = io.BytesIO()
    workbook.save(output)
    xlsx_bytes = output.getvalue()
    if len(xlsx_bytes) > MAX_XLSX_BYTES:
        raise ConversionError("Generated XLSX exceeds the 10 MB output limit.")
    load_workbook(io.BytesIO(xlsx_bytes), read_only=True).close()
    metadata = {"pages": page_count, "tables": table_count, "previews": previews}
    return xlsx_bytes, metadata


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input")
    parser.add_argument("--output")
    parser.add_argument("--base64-json", action="store_true")
    args = parser.parse_args()

    if args.base64_json:
        payload = json.load(__import__("sys").stdin)
        try:
            pdf_bytes = base64.b64decode(payload["pdf_base64"], validate=True)
            xlsx_bytes, metadata = convert_pdf_bytes(pdf_bytes)
            print(json.dumps({
                "ok": True,
                "xlsx_base64": base64.b64encode(xlsx_bytes).decode("ascii"),
                **metadata,
            }))
            return 0
        except Exception as exc:
            print(json.dumps({"ok": False, "error": str(exc)}))
            return 2

    if not args.input or not args.output:
        parser.error("--input and --output are required unless --base64-json is used")
    xlsx_bytes, metadata = convert_pdf_bytes(Path(args.input).read_bytes())
    Path(args.output).write_bytes(xlsx_bytes)
    print(json.dumps(metadata, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
