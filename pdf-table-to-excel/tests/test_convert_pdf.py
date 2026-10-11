import io
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
import convert_pdf  # noqa: E402
from convert_pdf import ConversionError, MAX_PDF_BYTES, convert_pdf_bytes  # noqa: E402


class ConvertPdfTests(unittest.TestCase):
    def make_pdf(self):
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4)
        table = Table([["Item", "Qty", "Price"], ["Rose", "2", "15.00"], ["Oud", "1", "40.00"]])
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ]))
        doc.build([table])
        return buf.getvalue()

    def test_real_pdf_table_becomes_valid_xlsx(self):
        xlsx, metadata = convert_pdf_bytes(self.make_pdf())
        self.assertEqual(metadata["tables"], 1)
        wb = load_workbook(io.BytesIO(xlsx), data_only=True)
        self.assertIn("Summary", wb.sheetnames)
        sheet = wb[metadata["previews"][0]["worksheet"]]
        self.assertEqual(sheet["A1"].value, "Item")
        self.assertEqual(sheet["A2"].value, "Rose")
        wb.close()

    def test_rejects_non_pdf(self):
        with self.assertRaisesRegex(ConversionError, "not a PDF"):
            convert_pdf_bytes(b"hello")

    def test_neutralizes_spreadsheet_formula_prefixes(self):
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4)
        dangerous = ['=HYPERLINK("https://example.com","click")', "+1+1", "-2+3", "@SUM(A1:A2)"]
        table = Table([["Value"], *[[value] for value in dangerous]])
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
        doc.build([table])
        xlsx, metadata = convert_pdf_bytes(buf.getvalue())
        wb = load_workbook(io.BytesIO(xlsx), data_only=False)
        sheet = wb[metadata["previews"][0]["worksheet"]]
        for row in range(2, 6):
            self.assertEqual(sheet.cell(row, 1).data_type, "s")
            self.assertTrue(sheet.cell(row, 1).value.startswith("'"))
        wb.close()

    def test_page_limit(self):
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4)
        from reportlab.platypus import PageBreak
        doc.build([Table([["A"]]), PageBreak(), Table([["B"]])])
        with patch.object(convert_pdf, "MAX_PDF_PAGES", 1):
            with self.assertRaisesRegex(ConversionError, "page limit"):
                convert_pdf_bytes(buf.getvalue())

    def test_cell_limit(self):
        with patch.object(convert_pdf, "MAX_TABLE_CELLS", 1):
            with self.assertRaisesRegex(ConversionError, "cell limit"):
                convert_pdf_bytes(self.make_pdf())

    def test_output_limit(self):
        with patch.object(convert_pdf, "MAX_XLSX_BYTES", 1):
            with self.assertRaisesRegex(ConversionError, "output limit"):
                convert_pdf_bytes(self.make_pdf())

    def test_rejects_oversize_pdf(self):
        with self.assertRaisesRegex(ConversionError, "20 MB"):
            convert_pdf_bytes(b"%PDF-" + b"0" * MAX_PDF_BYTES)


if __name__ == "__main__":
    unittest.main()
